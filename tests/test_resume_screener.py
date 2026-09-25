import io
import pytest
from fastapi.testclient import TestClient
from app import app
from modules.resume_screener import (
    extract_text_from_file_bytes,
    parse_resume_full,
    evaluate_ats_compatibility,
    generate_grounded_resume_verdict
)
from modules.qualifications import DEFAULT_JOB_TEMPLATES

client = TestClient(app)

SAMPLE_RESUME_TEXT = """
Alex Henderson
Email: alex.henderson@techdomain.io | Phone: (555) 234-5678
GitHub: github.com/alex-henderson | LinkedIn: linkedin.com/in/alex-henderson

PROFESSIONAL SUMMARY
Senior Backend Software Engineer with 6.5 years of experience architecting distributed microservices, scalable REST APIs, and event-driven data pipelines using Python, FastAPI, PostgreSQL, and Git.

TECHNICAL SKILLS
Languages & Frameworks: Python, FastAPI, SQL, PostgreSQL, REST API, Git, Docker, Redis
Cloud & DevOps: AWS, CI/CD Pipelines, Linux, Microservices

PROFESSIONAL EXPERIENCE
Lead Backend Engineer | CloudScale Systems (2020 - Present)
- Architected high-throughput REST API microservices in Python and FastAPI handling 45,000 req/s with 99.99% uptime.
- Optimized PostgreSQL relational queries and database indexing, reducing p99 query latency by 42%.
- Integrated automated CI/CD deployment pipelines using Docker and Git, accelerating release cycles by 35%.
- Mentored a squad of 5 junior backend engineers in code quality, type safety, and unit test automation.

Software Engineer | Apex Informatics (2018 - 2020)
- Developed modular Python API services and PostgreSQL persistence layers supporting 250,000 monthly active users.
- Refactored legacy monolithic backend into decoupled RESTful services, cutting memory consumption by 28%.

EDUCATION & CERTIFICATIONS
- B.Tech in Computer Science and Engineering | Apex Institute of Technology (2018)
- AWS Certified Solutions Architect - Associate
"""

def test_extract_text_plain_text():
    raw_bytes = SAMPLE_RESUME_TEXT.encode("utf-8")
    extracted = extract_text_from_file_bytes(raw_bytes, "resume.txt")
    assert "Alex Henderson" in extracted
    assert "FastAPI" in extracted

def test_parse_resume_full():
    parsed = parse_resume_full(SAMPLE_RESUME_TEXT)
    assert parsed["name"] == "Alex Henderson"
    assert parsed["email"] == "alex.henderson@techdomain.io"
    assert "(555) 234-5678" in parsed["phone"]
    assert len(parsed["links"]) >= 2
    assert "Python" in parsed["skills"]
    assert "PostgreSQL" in parsed["skills"]
    assert parsed["experience_years"] >= 5.0
    assert "B.Tech" in parsed["education"]
    assert "Experience" in parsed["sections_found"]
    assert "Skills" in parsed["sections_found"]
    assert "Education" in parsed["sections_found"]
    assert len(parsed["bullet_points"]) >= 4

def test_evaluate_ats_compatibility():
    parsed = parse_resume_full(SAMPLE_RESUME_TEXT)
    job = DEFAULT_JOB_TEMPLATES["JOB_SWE_01"]
    ats = evaluate_ats_compatibility(SAMPLE_RESUME_TEXT, parsed, job)

    assert ats["ats_score"] >= 75.0
    assert ats["section_score"] == 25.0
    assert ats["contact_score"] == 15.0
    assert ats["keyword_score"] >= 25.0
    assert ats["quantified_bullets_count"] >= 3
    assert ats["quantified_percentage"] > 0
    assert len(ats["checklist"]) >= 4

def test_generate_grounded_resume_verdict():
    parsed = parse_resume_full(SAMPLE_RESUME_TEXT)
    job = DEFAULT_JOB_TEMPLATES["JOB_SWE_01"]
    res = generate_grounded_resume_verdict(
        resume_text=SAMPLE_RESUME_TEXT,
        parsed_data=parsed,
        job_spec=job,
        mode="Demo Simulation Mode"
    )

    assert "candidate" in res
    assert "ats_analysis" in res
    assert "qualification_analysis" in res
    assert "ai_evaluation" in res
    assert res["ai_evaluation"]["decision"] in ["STRONG_HIRE", "HIRE", "INTERVIEW", "REJECT"]
    assert "grounded_feedback" in res
    assert len(res["grounded_feedback"]["strengths"]) >= 1
    assert "efs_assessment" in res
    assert res["efs_assessment"]["faithfulness_score"] >= 80.0
    assert "background_investigation" in res
    assert res["background_investigation"]["overall_status"] == "VERIFIED_FROM_PROVIDED_EVIDENCE"

def test_api_resume_screen_endpoint():
    payload = {
        "resume_text": SAMPLE_RESUME_TEXT,
        "job_id": "JOB_SWE_01",
        "mode": "Demo Simulation Mode"
    }
    res = client.post("/api/resume-screen", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert "ats_analysis" in data
    assert data["ats_analysis"]["ats_score"] >= 70.0
    assert "grounded_feedback" in data
    assert "strengths" in data["grounded_feedback"]
    assert "weaknesses" in data["grounded_feedback"]
    assert "recommendations" in data["grounded_feedback"]

def test_api_resume_upload_endpoint_txt():
    file_bytes = SAMPLE_RESUME_TEXT.encode("utf-8")
    files = {
        "file": ("alex_resume.txt", io.BytesIO(file_bytes), "text/plain")
    }
    data = {
        "job_id": "JOB_SWE_01",
        "mode": "Demo Simulation Mode"
    }
    res = client.post("/api/resume-upload", files=files, data=data)
    assert res.status_code == 200
    res_data = res.json()
    assert res_data["filename"] == "alex_resume.txt"
    assert "ats_analysis" in res_data
    assert res_data["ats_analysis"]["ats_score"] > 50.0
    assert "grounded_feedback" in res_data

def test_api_resume_upload_endpoint_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Alex Henderson", 0)
    doc.add_paragraph("Email: alex.henderson@techdomain.io | Phone: (555) 234-5678")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Python, FastAPI, SQL, PostgreSQL, REST API, Git, Docker")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Lead Backend Engineer with 6.5 years experience architecting REST microservices handling 45,000 req/s.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech in Computer Science")

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    docx_bytes = bio.read()

    # Test extraction function
    extracted = extract_text_from_file_bytes(docx_bytes, "resume.docx")
    assert "Alex Henderson" in extracted
    assert "FastAPI" in extracted

    # Test upload endpoint
    files = {
        "file": ("alex_resume.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }
    data = {
        "job_id": "JOB_SWE_01",
        "mode": "Demo Simulation Mode"
    }
    res = client.post("/api/resume-upload", files=files, data=data)
    assert res.status_code == 200
    res_data = res.json()
    assert res_data["filename"] == "alex_resume.docx"
    assert "ats_analysis" in res_data
    assert res_data["ats_analysis"]["ats_score"] > 50.0
    assert "grounded_feedback" in res_data


